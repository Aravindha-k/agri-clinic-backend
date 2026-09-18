"""Import operational villages from Excel: Village + village tamil name only."""

from __future__ import annotations

from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from masters.location_utils import normalize_village_name
from masters.models import Village

VILLAGE_HEADERS = {
    "village",
    "village name",
    "name",
}
TAMIL_HEADERS = {
    "village tamil name",
    "tamil name",
    "tamil_name",
    "name_ta",
    "village name ta",
    "tamil",
}
IGNORED_HEADERS = {
    "s no",
    "s.no",
    "s. no",
    "sno",
    "sl no",
    "sl.no",
    "field staff",
    "fieldstaff",
    "firka",
    "taluk",
    "district",
}


def _norm_header(value: object) -> str:
    return " ".join(str(value or "").strip().lower().split())


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


class Command(BaseCommand):
    help = (
        "Import operational villages from Excel using only Village and "
        "village tamil name. Ignores S NO / Field staff / Firka / Taluk / "
        "District. Does not assign employees."
    )

    def add_arguments(self, parser):
        parser.add_argument("path", type=str, help="Path to .xlsx workbook")
        parser.add_argument(
            "--dry-run",
            action="store_true",
            help="Parse and report without writing (default).",
        )
        parser.add_argument(
            "--execute",
            action="store_true",
            help="Write villages. Default is dry-run.",
        )

    def handle(self, *args, **options):
        path = Path(options["path"])
        if not path.exists():
            raise CommandError(f"File not found: {path}")

        dry_run = not options["execute"]
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            self._import_workbook(workbook, dry_run=dry_run)
        finally:
            workbook.close()

    def _import_workbook(self, workbook, *, dry_run: bool) -> None:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        try:
            header_row = next(rows)
        except StopIteration as exc:
            raise CommandError("Workbook is empty.") from exc

        village_idx = None
        tamil_idx = None
        for index, header in enumerate(header_row):
            key = _norm_header(header)
            if key in IGNORED_HEADERS:
                continue
            if village_idx is None and key in VILLAGE_HEADERS:
                village_idx = index
            elif tamil_idx is None and key in TAMIL_HEADERS:
                tamil_idx = index

        if village_idx is None:
            raise CommandError(
                "Required column 'Village' was not found. "
                "Expected a Village / Village Name column."
            )

        created = 0
        updated = 0
        skipped_blank = 0
        duplicates_in_file = 0
        index: dict[str, Village] = {}
        for village in Village.objects.only("id", "name", "name_ta"):
            key = normalize_village_name(village.name)
            index.setdefault(key, village)

        seen_in_file: set[str] = set()
        to_create: list[Village] = []
        update_map: dict[int, Village] = {}

        for row in rows:
            if not row:
                continue
            name = _cell_text(row[village_idx] if village_idx < len(row) else "")
            tamil = ""
            if tamil_idx is not None and tamil_idx < len(row):
                tamil = _cell_text(row[tamil_idx])
            if not name:
                skipped_blank += 1
                continue
            key = normalize_village_name(name)
            if key in seen_in_file:
                duplicates_in_file += 1
                existing = index.get(key)
                if existing and tamil and not existing.name_ta:
                    existing.name_ta = tamil
                    if existing.pk:
                        update_map[existing.pk] = existing
                continue
            seen_in_file.add(key)

            existing = index.get(key)
            if existing:
                changed = False
                if tamil and existing.name_ta != tamil:
                    existing.name_ta = tamil
                    changed = True
                if not existing.is_active:
                    existing.is_active = True
                    changed = True
                if changed and existing.pk:
                    update_map[existing.pk] = existing
                    updated += 1
            else:
                village = Village(
                    name=name,
                    name_ta=tamil,
                    is_active=True,
                )
                to_create.append(village)
                index[key] = village
                created += 1

        if dry_run:
            self.stdout.write(
                self.style.WARNING(
                    f"DRY RUN created={created} updated={updated} "
                    f"blank={skipped_blank} file_dupes={duplicates_in_file}"
                )
            )
            return

        with transaction.atomic():
            if to_create:
                Village.objects.bulk_create(to_create, batch_size=500)
            if update_map:
                Village.objects.bulk_update(
                    list(update_map.values()), ["name_ta", "is_active"], batch_size=500
                )

        self.stdout.write(
            self.style.SUCCESS(
                f"Imported villages created={created} updated={updated} "
                f"blank={skipped_blank} file_dupes={duplicates_in_file}. "
                "No employee assignments were created."
            )
        )
