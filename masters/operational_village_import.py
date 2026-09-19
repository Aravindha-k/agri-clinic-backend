"""
Plan and apply operational Village + EmployeeLocationAssignment import.

Village-only. No District/Taluk/Firka. Default is dry-run (zero writes).
"""

from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path

from django.db import transaction
from openpyxl import load_workbook

from accounts.models import EmployeeLocationAssignment, EmployeeProfile
from masters.location_utils import normalize_village_name
from masters.models import Village

VILLAGE_HEADERS = {
    "village",
    "village name",
    "name",
}
TAMIL_HEADERS = {
    "village tamil name",
    "tamil name",
    "tamil_name",
    "name_ta",
    "village name ta",
    "tamil",
}
STAFF_HEADERS = {
    "field staff",
    "fieldstaff",
    "field_staff",
    "staff",
    "employee",
    "employee name",
}
IGNORED_HEADERS = {
    "s no",
    "s.no",
    "s. no",
    "sno",
    "sl no",
    "sl.no",
    "firka",
    "taluk",
    "district",
}

# Production Field staff -> EmployeeProfile.employee_id. Verified against DB.
EXPECTED_EMPLOYEE_IDS = {
    "kaviyarasan": "KAC-0003",
    "sasikumar": "KAC-0004",
    "sathish": "KAC-0005",
    "selvamani": "KAC-0006",
}


def _norm_header(value: object) -> str:
    return " ".join(str(value or "").strip().lower().split())


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def normalize_person_name(name: str) -> str:
    return " ".join((name or "").strip().split()).casefold()


@dataclass
class TamilConflict:
    village_key: str
    display_name: str
    tamil_values: list[str]
    excel_rows: list[int]


@dataclass
class EmployeeMatch:
    staff_name: str
    status: str  # found | not_found | ambiguous
    profile: EmployeeProfile | None = None
    candidates: list[EmployeeProfile] = field(default_factory=list)


@dataclass
class EmployeeBreakdown:
    staff_name: str
    employee_id_code: str
    profile_pk: int | None
    excel_rows: int = 0
    unique_villages: int = 0
    assignments_to_create: int = 0
    assignments_to_reuse: int = 0


@dataclass
class ImportPlan:
    total_data_rows: int = 0
    total_valid_rows: int = 0
    skipped_rows: list[str] = field(default_factory=list)
    invalid_missing_village: list[str] = field(default_factory=list)
    unique_villages: int = 0
    villages_to_create: int = 0
    villages_to_reuse: int = 0
    tamil_present: int = 0
    tamil_blank: int = 0
    blank_tamil_villages: list[str] = field(default_factory=list)
    tamil_conflicts: list[TamilConflict] = field(default_factory=list)
    employees_found: list[str] = field(default_factory=list)
    employees_not_found: list[str] = field(default_factory=list)
    ambiguous_employees: list[str] = field(default_factory=list)
    employee_id_mismatches: list[str] = field(default_factory=list)
    employee_matches: dict[str, EmployeeMatch] = field(default_factory=dict)
    assignments_to_create: int = 0
    assignments_to_reuse: int = 0
    unique_employee_village_pairs: int = 0
    duplicate_excel_village_rows: int = 0
    duplicate_employee_village_rows: int = 0
    shared_villages: dict[str, list[str]] = field(default_factory=dict)
    per_employee: dict[str, EmployeeBreakdown] = field(default_factory=dict)
    # Internal apply payloads
    village_create: dict[str, dict[str, str]] = field(default_factory=dict)
    village_update_ta: dict[str, str] = field(default_factory=dict)
    assignment_keys: set[tuple[int, str]] = field(default_factory=set)
    blocking_errors: list[str] = field(default_factory=list)


def _resolve_columns(header_row) -> tuple[int | None, int | None, int | None]:
    village_idx = None
    tamil_idx = None
    staff_idx = None
    for index, header in enumerate(header_row):
        key = _norm_header(header)
        if key in IGNORED_HEADERS:
            continue
        if village_idx is None and key in VILLAGE_HEADERS:
            village_idx = index
        elif tamil_idx is None and key in TAMIL_HEADERS:
            tamil_idx = index
        elif staff_idx is None and key in STAFF_HEADERS:
            staff_idx = index
    return village_idx, tamil_idx, staff_idx


def _employee_name_keys(profile: EmployeeProfile) -> set[str]:
    user = profile.user
    keys: set[str] = set()
    for raw in (
        user.first_name,
        user.last_name,
        user.get_full_name(),
        user.username,
    ):
        key = normalize_person_name(raw or "")
        if key:
            keys.add(key)
    # "First Last" also try first token only when full name has spaces
    full = normalize_person_name(user.get_full_name() or "")
    if " " in full:
        keys.add(full.split(" ", 1)[0])
    return keys


def build_employee_index() -> dict[str, list[EmployeeProfile]]:
    index: dict[str, list[EmployeeProfile]] = defaultdict(list)
    seen_pk: dict[str, set[int]] = defaultdict(set)
    qs = EmployeeProfile.objects.select_related("user").all()
    for profile in qs:
        for key in _employee_name_keys(profile):
            if profile.pk in seen_pk[key]:
                continue
            seen_pk[key].add(profile.pk)
            index[key].append(profile)
    return index


def match_field_staff(
    staff_name: str, index: dict[str, list[EmployeeProfile]]
) -> EmployeeMatch:
    key = normalize_person_name(staff_name)
    if not key:
        return EmployeeMatch(staff_name=staff_name, status="not_found")
    candidates = index.get(key, [])
    if len(candidates) == 1:
        return EmployeeMatch(
            staff_name=staff_name, status="found", profile=candidates[0]
        )
    if len(candidates) > 1:
        return EmployeeMatch(
            staff_name=staff_name,
            status="ambiguous",
            candidates=list(candidates),
        )
    return EmployeeMatch(staff_name=staff_name, status="not_found")


def collect_plan(path: Path) -> ImportPlan:
    """Read-only inventory. Wrapped so accidental writes cannot persist."""
    with transaction.atomic():
        plan = _build_plan(path)
        transaction.set_rollback(True)
        return plan


def _build_plan(path: Path) -> ImportPlan:
    plan = ImportPlan()
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows_iter = sheet.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter)
        except StopIteration as exc:
            raise ValueError("Workbook is empty.") from exc

        village_idx, tamil_idx, staff_idx = _resolve_columns(header_row)
        if village_idx is None:
            raise ValueError(
                "Required column 'Village' was not found. "
                "Expected a Village / Village Name column."
            )
        if staff_idx is None:
            raise ValueError(
                "Required column 'Field staff' was not found for territory import."
            )

        # Existing villages by normalized name
        village_db: dict[str, Village] = {}
        for village in Village.objects.only("id", "name", "name_ta", "is_active"):
            village_db.setdefault(normalize_village_name(village.name), village)

        # Existing operational assignments: (employee_pk, village_key)
        existing_assignments: set[tuple[int, str]] = set()
        for row in (
            EmployeeLocationAssignment.objects.filter(
                is_operational=True, village__isnull=False
            )
            .select_related("village")
            .only("employee_id", "village__name")
        ):
            existing_assignments.add(
                (
                    row.employee_id,
                    normalize_village_name(row.village.name),
                )
            )

        employee_index = build_employee_index()
        staff_cache: dict[str, EmployeeMatch] = {}

        # village_key -> {tamil_values set, display_name, excel_rows, staff names}
        village_meta: dict[str, dict] = {}
        # (emp_pk or staff_key, village_key) occurrence count in file
        pair_counts: dict[tuple[str, str], int] = defaultdict(int)
        # staff_key -> set of village keys (valid rows only)
        staff_villages: dict[str, set[str]] = defaultdict(set)
        staff_row_counts: dict[str, int] = defaultdict(int)

        excel_row_num = 1  # header is row 1
        for row in rows_iter:
            excel_row_num += 1
            if not row:
                continue
            name = _cell_text(row[village_idx] if village_idx < len(row) else "")
            tamil = ""
            if tamil_idx is not None and tamil_idx < len(row):
                tamil = _cell_text(row[tamil_idx])
            staff = _cell_text(row[staff_idx] if staff_idx < len(row) else "")

            if not name and not staff and not tamil:
                continue

            plan.total_data_rows += 1

            if not name:
                detail = (
                    f"INVALID_MISSING_VILLAGE: Excel row {excel_row_num} | "
                    f"Field staff={staff!r} | village tamil name={tamil!r}"
                )
                plan.invalid_missing_village.append(detail)
                plan.skipped_rows.append(detail)
                continue
            if not staff:
                detail = (
                    f"INVALID_MISSING_FIELD_STAFF: Excel row {excel_row_num} | "
                    f"Village={name!r} | village tamil name={tamil!r}"
                )
                plan.skipped_rows.append(detail)
                continue

            plan.total_valid_rows += 1

            vkey = normalize_village_name(name)
            meta = village_meta.setdefault(
                vkey,
                {
                    "display_name": name,
                    "tamils": set(),
                    "tamil_rows": defaultdict(list),
                    "rows": [],
                    "staff": set(),
                },
            )
            # Prefer first non-empty display casing as canonical create name
            if not meta["display_name"]:
                meta["display_name"] = name
            meta["rows"].append(excel_row_num)
            if tamil:
                meta["tamils"].add(tamil)
                meta["tamil_rows"][tamil].append(excel_row_num)

            staff_key = normalize_person_name(staff)
            if staff_key not in staff_cache:
                staff_cache[staff_key] = match_field_staff(staff, employee_index)
            match = staff_cache[staff_key]
            # Keep original display staff name from first sighting
            if not match.staff_name:
                match.staff_name = staff

            pair_key = (staff_key, vkey)
            pair_counts[pair_key] += 1
            if pair_counts[pair_key] == 1:
                staff_villages[staff_key].add(vkey)
                meta["staff"].add(staff_key)
            staff_row_counts[staff_key] += 1

        # Village stats + tamil conflicts
        # blank + nonblank Tamil => use nonblank (blanks never enter meta["tamils"])
        for vkey, meta in village_meta.items():
            plan.unique_villages += 1
            tamils = sorted(meta["tamils"])
            chosen_ta = ""
            if len(tamils) > 1:
                conflict_rows: list[int] = []
                for ta in tamils:
                    conflict_rows.extend(meta["tamil_rows"].get(ta, []))
                plan.tamil_conflicts.append(
                    TamilConflict(
                        village_key=vkey,
                        display_name=meta["display_name"],
                        tamil_values=tamils,
                        excel_rows=sorted(set(conflict_rows)) or list(meta["rows"]),
                    )
                )
            elif len(tamils) == 1:
                plan.tamil_present += 1
                chosen_ta = tamils[0]
            else:
                plan.tamil_blank += 1
                plan.blank_tamil_villages.append(meta["display_name"])

            existing = village_db.get(vkey)
            if existing:
                plan.villages_to_reuse += 1
                if chosen_ta and not (existing.name_ta or "").strip():
                    plan.village_update_ta[vkey] = chosen_ta
                elif (
                    chosen_ta
                    and (existing.name_ta or "").strip()
                    and normalize_person_name(existing.name_ta)
                    != normalize_person_name(chosen_ta)
                    and chosen_ta not in {(existing.name_ta or "").strip()}
                ):
                    # DB has different tamil than Excel nonblank — treat as conflict
                    if not any(c.village_key == vkey for c in plan.tamil_conflicts):
                        plan.tamil_conflicts.append(
                            TamilConflict(
                                village_key=vkey,
                                display_name=meta["display_name"],
                                tamil_values=sorted(
                                    {
                                        (existing.name_ta or "").strip(),
                                        chosen_ta,
                                    }
                                ),
                                excel_rows=list(meta["rows"]),
                            )
                        )
            else:
                plan.villages_to_create += 1
                plan.village_create[vkey] = {
                    "name": meta["display_name"],
                    "name_ta": chosen_ta,
                }

            # Duplicate village rows = total rows for village beyond first
            if len(meta["rows"]) > 1:
                plan.duplicate_excel_village_rows += len(meta["rows"]) - 1

        # Employee match summary + required production ID verification
        for staff_key, match in staff_cache.items():
            display = match.staff_name or staff_key
            plan.employee_matches[display] = match
            if match.status == "found":
                plan.employees_found.append(display)
                profile = match.profile
                assert profile is not None
                expected_code = EXPECTED_EMPLOYEE_IDS.get(staff_key)
                if expected_code and profile.employee_id != expected_code:
                    plan.employee_id_mismatches.append(
                        f"{display}: DB employee_id={profile.employee_id!r} "
                        f"expected={expected_code!r} pk={profile.pk}"
                    )
                bd = EmployeeBreakdown(
                    staff_name=display,
                    employee_id_code=profile.employee_id,
                    profile_pk=profile.pk,
                    excel_rows=staff_row_counts.get(staff_key, 0),
                    unique_villages=len(staff_villages.get(staff_key, set())),
                )
                for vkey in staff_villages.get(staff_key, set()):
                    pair = (profile.pk, vkey)
                    if pair in existing_assignments:
                        plan.assignments_to_reuse += 1
                        bd.assignments_to_reuse += 1
                    else:
                        plan.assignments_to_create += 1
                        bd.assignments_to_create += 1
                        plan.assignment_keys.add(pair)
                plan.per_employee[display] = bd
            elif match.status == "ambiguous":
                plan.ambiguous_employees.append(display)
                plan.per_employee[display] = EmployeeBreakdown(
                    staff_name=display,
                    employee_id_code="",
                    profile_pk=None,
                    excel_rows=staff_row_counts.get(staff_key, 0),
                    unique_villages=len(staff_villages.get(staff_key, set())),
                )
            else:
                plan.employees_not_found.append(display)
                plan.per_employee[display] = EmployeeBreakdown(
                    staff_name=display,
                    employee_id_code="",
                    profile_pk=None,
                    excel_rows=staff_row_counts.get(staff_key, 0),
                    unique_villages=len(staff_villages.get(staff_key, set())),
                )

        # Duplicate employee/village rows in Excel
        for (_staff, _vkey), count in pair_counts.items():
            if count > 1:
                plan.duplicate_employee_village_rows += count - 1

        plan.unique_employee_village_pairs = len(pair_counts)

        # Shared villages: >1 distinct staff in Excel
        for vkey, meta in village_meta.items():
            staff_keys = sorted(meta["staff"])
            if len(staff_keys) > 1:
                names = []
                for sk in staff_keys:
                    m = staff_cache.get(sk)
                    names.append(m.staff_name if m else sk)
                plan.shared_villages[meta["display_name"]] = names

        if plan.employees_not_found:
            plan.blocking_errors.append("EMPLOYEE_NOT_FOUND")
        if plan.ambiguous_employees:
            plan.blocking_errors.append("AMBIGUOUS_EMPLOYEE")
        if plan.employee_id_mismatches:
            plan.blocking_errors.append("EMPLOYEE_ID_MISMATCH")
        if plan.tamil_conflicts:
            plan.blocking_errors.append("TAMIL_NAME_CONFLICT")

        return plan
    finally:
        workbook.close()


def execute_plan(plan: ImportPlan) -> dict[str, int]:
    """
    Apply village creates/updates and assignment upserts.
    Refuses when plan.blocking_errors is non-empty.
    """
    if plan.blocking_errors:
        raise ValueError(
            "Refusing execute due to blocking errors: "
            + ", ".join(plan.blocking_errors)
        )

    created_villages = 0
    updated_villages = 0
    created_assignments = 0

    with transaction.atomic():
        # Build/refresh village index
        village_by_key: dict[str, Village] = {}
        for village in Village.objects.all():
            village_by_key.setdefault(normalize_village_name(village.name), village)

        for vkey, payload in plan.village_create.items():
            if vkey in village_by_key:
                continue
            village = Village.objects.create(
                name=payload["name"],
                name_ta=payload.get("name_ta") or "",
                is_active=True,
                district=None,
                taluk=None,
            )
            village_by_key[vkey] = village
            created_villages += 1

        for vkey, tamil in plan.village_update_ta.items():
            village = village_by_key.get(vkey)
            if village is None:
                continue
            if not (village.name_ta or "").strip() and tamil:
                village.name_ta = tamil
                village.is_active = True
                village.save(update_fields=["name_ta", "is_active", "updated_at"])
                updated_villages += 1

        for emp_pk, vkey in plan.assignment_keys:
            village = village_by_key.get(vkey)
            if village is None:
                raise ValueError(f"Missing village for assignment key {vkey}")
            obj = (
                EmployeeLocationAssignment.objects.filter(
                    employee_id=emp_pk,
                    village=village,
                    is_operational=True,
                )
                .order_by("id")
                .first()
            )
            if obj is None:
                EmployeeLocationAssignment.objects.create(
                    employee_id=emp_pk,
                    village=village,
                    district=None,
                    taluk=None,
                    is_active=True,
                    is_operational=True,
                )
                created_assignments += 1
            else:
                changed = False
                if obj.district_id is not None:
                    obj.district = None
                    changed = True
                if obj.taluk_id is not None:
                    obj.taluk = None
                    changed = True
                if not obj.is_active:
                    obj.is_active = True
                    changed = True
                if changed:
                    obj.save(
                        update_fields=[
                            "district",
                            "taluk",
                            "is_active",
                            "updated_at",
                        ]
                    )

    return {
        "villages_created": created_villages,
        "villages_updated": updated_villages,
        "assignments_created": created_assignments,
    }


def format_plan(plan: ImportPlan, *, dry_run: bool) -> str:
    lines: list[str] = []
    mode = "DRY RUN — no rows will be written." if dry_run else "EXECUTE — writing villages + assignments."
    lines.append(mode)
    lines.append("")
    lines.append(f"TOTAL EXCEL ROWS: {plan.total_data_rows}")
    lines.append(f"TOTAL VALID EXCEL ROWS: {plan.total_valid_rows}")
    lines.append(f"INVALID MISSING VILLAGE ROWS: {len(plan.invalid_missing_village)}")
    for item in plan.invalid_missing_village:
        lines.append(f"  - {item}")
    lines.append(f"UNIQUE VILLAGES: {plan.unique_villages}")
    lines.append(f"VILLAGES TO CREATE: {plan.villages_to_create}")
    lines.append(f"VILLAGES TO REUSE: {plan.villages_to_reuse}")
    lines.append(f"TAMIL NAMES PRESENT: {plan.tamil_present}")
    lines.append(f"TAMIL NAMES BLANK: {plan.tamil_blank}")
    lines.append(f"TAMIL NAME CONFLICTS: {len(plan.tamil_conflicts)}")
    if plan.tamil_conflicts:
        for c in plan.tamil_conflicts:
            lines.append(
                f"  TAMIL_NAME_CONFLICT: {c.display_name} | "
                f"values={c.tamil_values} | rows={c.excel_rows}"
            )
    if plan.blank_tamil_villages:
        lines.append("BLANK TAMIL VILLAGES (sample up to 30):")
        for name in plan.blank_tamil_villages[:30]:
            lines.append(f"  - {name}")
        if len(plan.blank_tamil_villages) > 30:
            lines.append(f"  ... +{len(plan.blank_tamil_villages) - 30} more")

    lines.append("")
    lines.append(f"FIELD STAFF FOUND: {len(plan.employees_found)}")
    lines.append("EMPLOYEE MATCHES:")
    for name, match in sorted(plan.employee_matches.items(), key=lambda x: x[0].casefold()):
        if match.status == "found" and match.profile is not None:
            p = match.profile
            lines.append(
                f"  {name} -> EmployeeProfile pk={p.pk} code={p.employee_id} "
                f"user={p.user.username}"
            )
        elif match.status == "ambiguous":
            ids = ", ".join(
                f"pk={c.pk}/code={c.employee_id}" for c in match.candidates
            )
            lines.append(f"  {name} -> AMBIGUOUS_EMPLOYEE [{ids}]")
        else:
            lines.append(f"  {name} -> EMPLOYEE_NOT_FOUND")

    lines.append(f"EMPLOYEE NOT FOUND: {plan.employees_not_found or 'none'}")
    lines.append(f"AMBIGUOUS EMPLOYEES: {plan.ambiguous_employees or 'none'}")
    if plan.employee_id_mismatches:
        lines.append("EMPLOYEE_ID_MISMATCH:")
        for item in plan.employee_id_mismatches:
            lines.append(f"  - {item}")
    lines.append("")

    preferred_order = [
        "Kaviyarasan",
        "Sasikumar",
        "Sathish",
        "Selvamani",
    ]
    ordered = []
    remaining = dict(plan.per_employee)
    for label in preferred_order:
        key = next(
            (k for k in remaining if normalize_person_name(k) == normalize_person_name(label)),
            None,
        )
        if key is not None:
            ordered.append(key)
            remaining.pop(key)
    ordered.extend(sorted(remaining.keys(), key=str.casefold))

    for label in ordered:
        bd = plan.per_employee[label]
        lines.append(f"{label}:")
        lines.append(f"  EmployeeProfile: pk={bd.profile_pk}")
        lines.append(f"  Employee ID: {bd.employee_id_code or 'n/a'}")
        lines.append(f"  Excel rows: {bd.excel_rows}")
        lines.append(f"  Unique villages: {bd.unique_villages}")
        lines.append(f"  Assignments to create: {bd.assignments_to_create}")
        lines.append(f"  Assignments to reuse: {bd.assignments_to_reuse}")
        lines.append("")

    lines.append(f"SHARED VILLAGES: {len(plan.shared_villages)}")
    for vname, staff in sorted(plan.shared_villages.items(), key=lambda x: x[0].casefold()):
        lines.append(f"  {vname}: {', '.join(staff)}")

    lines.append("")
    lines.append(f"UNIQUE EMPLOYEE/VILLAGE PAIRS: {plan.unique_employee_village_pairs}")
    lines.append(f"DUPLICATE VILLAGE ROWS: {plan.duplicate_excel_village_rows}")
    lines.append(
        f"DUPLICATE EMPLOYEE/VILLAGE ROWS: {plan.duplicate_employee_village_rows}"
    )
    lines.append(f"ASSIGNMENTS TO CREATE: {plan.assignments_to_create}")
    lines.append(f"ASSIGNMENTS TO REUSE: {plan.assignments_to_reuse}")
    lines.append(f"SKIPPED/INVALID ROWS: {len(plan.skipped_rows)}")
    for item in plan.skipped_rows[:40]:
        lines.append(f"  - {item}")
    if len(plan.skipped_rows) > 40:
        lines.append(f"  ... +{len(plan.skipped_rows) - 40} more")

    lines.append("")
    if plan.blocking_errors:
        lines.append("BLOCKING ERRORS: " + ", ".join(plan.blocking_errors))
        lines.append("Execute is refused until these are reviewed/resolved.")
    else:
        lines.append("BLOCKING ERRORS: none")

    if dry_run:
        lines.append("DRY RUN ZERO WRITES: YES")
        lines.append("PRODUCTION IMPORT EXECUTED: NO")
    else:
        lines.append("DRY RUN ZERO WRITES: NO")
    return "\n".join(lines)
